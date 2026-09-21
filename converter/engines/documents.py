"""
Document and Text Conversion Engine for OmniConvert.
Converts PDF, DOCX, TXT, MD, HTML, and rich document formats with vector rendering and text preservation.
"""

import os
import sys
import time
from typing import List, Dict, Any, Optional
from converter.engines.base import BaseConverter, ConversionProgress, ProgressCallback, CancellationToken

# Document libraries
try:
    import pypdf
except ImportError:
    pypdf = None

try:
    import docx
except ImportError:
    docx = None


class DocumentConverter(BaseConverter):
    """Universal document and rich text converter."""

    DOC_FORMATS = ["pdf", "docx", "txt", "md", "html", "htm", "rtf", "json", "yaml", "yml", "xml"]

    @property
    def name(self) -> str:
        return "Universal Document & Rich Text Engine"

    @property
    def supported_source_formats(self) -> List[str]:
        return self.DOC_FORMATS

    @property
    def supported_target_formats(self) -> List[str]:
        return self.DOC_FORMATS

    def can_convert(self, src_ext: str, dst_ext: str) -> bool:
        src = src_ext.lower().lstrip(".")
        dst = dst_ext.lower().lstrip(".")
        if src == dst:
            return False
        return src in self.DOC_FORMATS and dst in self.DOC_FORMATS

    def _read_docx_text(self, path: str) -> str:
        if not docx:
            raise RuntimeError("python-docx is required for DOCX conversion.")
        doc = docx.Document(path)
        full_text = []
        for p in doc.paragraphs:
            full_text.append(p.text)
        for table in doc.tables:
            for row in table.rows:
                full_text.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(full_text)

    def _read_pdf_text(self, path: str) -> str:
        if not pypdf:
            raise RuntimeError("pypdf is required for PDF text extraction.")
        reader = pypdf.PdfReader(path)
        pages_text = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            pages_text.append(f"--- Page {i+1} ---\n{text}")
        return "\n\n".join(pages_text)

    def _render_text_to_pdf(self, text_or_html: str, output_path: str, is_html: bool = False):
        """Render text or HTML to a crisp vector PDF using PyQt6's QPdfWriter & QTextDocument."""
        from PyQt6.QtGui import QTextDocument, QPdfWriter, QPageSize, QPageLayout
        from PyQt6.QtCore import QSizeF, QMarginsF

        doc = QTextDocument()
        if is_html:
            doc.setHtml(text_or_html)
        else:
            # Escape plain text to safe HTML preserving line breaks
            html_content = (
                text_or_html.replace("&", "&amp;")
                .replace("<", "&lt;")
                .replace(">", "&gt;")
                .replace("\n", "<br>")
            )
            doc.setHtml(f"<div style='font-family: Arial, sans-serif; font-size: 11pt; line-height: 1.5;'>{html_content}</div>")

        writer = QPdfWriter(output_path)
        writer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
        layout = QPageLayout(
            QPageSize(QPageSize.PageSizeId.A4),
            QPageLayout.Orientation.Portrait,
            QMarginsF(15, 15, 15, 15),
            QPageLayout.Unit.Millimeter
        )
        writer.setPageLayout(layout)
        doc.print(writer)

    def convert(
        self,
        source_path: str,
        target_path: str,
        options: Optional[Dict[str, Any]] = None,
        progress_callback: Optional[ProgressCallback] = None,
        cancel_token: Optional[CancellationToken] = None,
    ) -> bool:
        options = options or {}
        src_ext = os.path.splitext(source_path)[1].lower().lstrip(".")
        dst_ext = os.path.splitext(target_path)[1].lower().lstrip(".")
        start_time = time.time()
        source_size = os.path.getsize(source_path) if os.path.exists(source_path) else 0

        if progress_callback:
            progress_callback(
                ConversionProgress(
                    progress=15.0,
                    status_text=f"Reading {src_ext.upper()} document...",
                    current_bytes=0,
                    total_bytes=source_size
                )
            )

        if cancel_token and cancel_token.is_cancelled:
            raise RuntimeError("Conversion cancelled by user.")

        # Step 1: Extract intermediate content
        raw_text = ""
        html_content = ""

        if src_ext in ("txt", "json", "yaml", "yml", "xml", "rtf"):
            with open(source_path, "r", encoding="utf-8", errors="replace") as f:
                raw_text = f.read()
        elif src_ext == "md":
            with open(source_path, "r", encoding="utf-8", errors="replace") as f:
                raw_text = f.read()
            # Basic markdown to HTML formatting
            try:
                import markdown2
                html_content = markdown2.markdown(raw_text)
            except ImportError:
                # Simple fallback
                html_content = f"<pre style='font-family: monospace;'>{raw_text}</pre>"
        elif src_ext in ("html", "htm"):
            with open(source_path, "r", encoding="utf-8", errors="replace") as f:
                html_content = f.read()
            raw_text = html_content  # Fallback text
        elif src_ext == "docx":
            raw_text = self._read_docx_text(source_path)
        elif src_ext == "pdf":
            raw_text = self._read_pdf_text(source_path)

        if progress_callback:
            progress_callback(
                ConversionProgress(
                    progress=50.0,
                    status_text=f"Generating {dst_ext.upper()} output...",
                    current_bytes=source_size // 2,
                    total_bytes=source_size
                )
            )

        if cancel_token and cancel_token.is_cancelled:
            raise RuntimeError("Conversion cancelled by user.")

        # Step 2: Write target format
        if dst_ext == "pdf":
            content_to_render = html_content if html_content else raw_text
            self._render_text_to_pdf(content_to_render, target_path, is_html=bool(html_content))

        elif dst_ext == "docx":
            if not docx:
                raise RuntimeError("python-docx is required for DOCX generation.")
            doc = docx.Document()
            # Add content lines
            for line in raw_text.splitlines():
                doc.add_paragraph(line)
            doc.save(target_path)

        elif dst_ext == "html":
            if html_content:
                final_html = html_content
            else:
                body = raw_text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br>")
                final_html = f"<!DOCTYPE html><html><head><meta charset='utf-8'><title>Converted Document</title></head><body>{body}</body></html>"
            with open(target_path, "w", encoding="utf-8") as f:
                f.write(final_html)

        elif dst_ext in ("txt", "md", "json", "yaml", "yml", "xml", "rtf"):
            with open(target_path, "w", encoding="utf-8") as f:
                f.write(raw_text)

        elapsed = max(0.001, time.time() - start_time)
        final_size = os.path.getsize(target_path) if os.path.exists(target_path) else source_size

        if progress_callback:
            progress_callback(
                ConversionProgress(
                    progress=100.0,
                    speed_bytes_sec=final_size / elapsed,
                    eta_seconds=0.0,
                    status_text="Document converted successfully",
                    current_bytes=final_size,
                    total_bytes=final_size
                )
            )

        return True
